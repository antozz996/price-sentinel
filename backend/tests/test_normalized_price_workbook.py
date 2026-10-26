from io import BytesIO
import hashlib

from openpyxl import Workbook

from app.services.normalized_price_workbook import (
    EXPECTED_TABS,
    NEW_FORMAT_TABS,
    NEW_FORMAT_HEADERS,
    REQUIRED_HEADERS,
    parse_normalized_price_workbook,
)


HEADERS = list(REQUIRED_HEADERS.values())


def make_workbook() -> bytes:
    workbook = Workbook()
    first = workbook.active
    first.title = "FOOD"
    for tab in EXPECTED_TABS:
        sheet = first if tab == "FOOD" else workbook.create_sheet(tab)
        sheet.append([f"{tab} — PREZZI ATTUALI"])
        sheet.append(["Prezzi attuali approvati"])
        sheet.append(["Aggiornamenti come nuove righe"])
        sheet.append(HEADERS)
    workbook["FOOD"].append([
        "Carni", "Entrecôte", "MARR", "ENTRECOTE PRUSSIANA", "A123", "Marca A",
        "Confezione 2 kg", "€ 30,00", 1, "kg", "15,00", "Prezzo attuale approvato",
        "FOOD!C2", None, "mensile",
    ])
    workbook["FOOD"].append([
        "Carni", "Pollo", "DAC", "Pollo intero", "", "", "1 conf", "€ 8,00", 1,
        "", "8,00", "Prezzo attuale approvato", "FOOD!D3", None, "mensile",
    ])
    workbook["FOOD"].append([
        "Carni", "Duplicato", "DAC", "Duplicato A", "", "", "1 kg", "€ 2,00", 1,
        "kg", "2,00", "Prezzo attuale approvato", "FOOD!D4", None, "mensile",
    ])
    workbook["FOOD"].append([
        "Carni", "Duplicato", "DAC", "Duplicato B", "", "", "2 kg", "€ 3,00", 1,
        "kg", "3,00", "Prezzo attuale approvato", "FOOD!D5", None, "mensile",
    ])
    workbook["FOOD"].append([
        "Carni", "Ignorato", "MARR", "", "", "", "", "", "", "", "",
        "Duplicata - non importare", "FOOD!C6", None, "",
    ])
    workbook["BEVERAGE"].append([
        "Bevande", "Succo monodose", "MELIUS", "Succo 1 pezzo", "", "", "1 pezzo",
        "€ 6,20", 1, "pz", None, "Prezzo attuale approvato", "BEVERAGE!D2", None, "mensile",
    ])
    workbook["BEVERAGE"].append([
        "Bevande", "Prodotto futuro", "MELIUS", "", "", "", "", None, None, None,
        None, "Prezzo attuale approvato", "BEVERAGE!D3", None, "mensile",
    ])
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def test_normalized_workbook_preview_keeps_only_approved_offers_and_blocks_unclear_rows() -> None:
    report = parse_normalized_price_workbook(make_workbook())

    assert report["mode"] == "dry_run_only"
    assert report["no_contract_writes"] is True
    assert report["candidate_count"] == 5
    assert report["excluded_count"] == 1
    assert report["skipped_empty_count"] == 1
    offers = report["candidates"]
    entrecote = next(item for item in offers if item["product"] == "Entrecôte")
    assert entrecote["numeric_price"] == "30.00"
    assert entrecote["source_cell"] == "H5"
    assert entrecote["reference_unit"] == "kg"
    assert entrecote["packaging_note"] is not None
    assert entrecote["status"] == "needs_mapping"

    single_unit = next(item for item in offers if item["product"] == "Succo monodose")
    assert single_unit["quoted_value"] == "€ 6,20"
    assert single_unit["numeric_price"] == "6.20"
    assert single_unit["source_cell"] == "H5"
    assert single_unit["status"] == "needs_mapping"

    chicken = next(item for item in offers if item["product"] == "Pollo")
    assert chicken["status"] == "blocked"
    assert "unit_not_specified" in chicken["reasons"]

    duplicate_rows = [item for item in offers if item["product"] == "Duplicato"]
    assert len(duplicate_rows) == 2
    assert all(item["status"] == "blocked" for item in duplicate_rows)
    assert all("duplicate_supplier_product_offer" in item["reasons"] for item in duplicate_rows)
    assert all(item["product"] != "Prodotto futuro" for item in offers)


def test_package_price_is_not_converted_when_declared_quantity_is_not_one() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "FOOD"
    for tab in EXPECTED_TABS:
        sheet = first if tab == "FOOD" else workbook.create_sheet(tab)
        sheet.append([f"{tab} — PREZZI ATTUALI"])
        sheet.append(["Prezzi attuali approvati"])
        sheet.append(["Aggiornamenti come nuove righe"])
        sheet.append(HEADERS)
    workbook["FOOD"].append([
        "Carni", "Gelato", "MELIUS", "Gelato 2,8 kg", "", "", "Confezione 2,8 kg",
        "€ 14,00", "2,8", "kg", None, "Prezzo attuale approvato", "FOOD!D2", None, "mensile",
    ])
    buffer = BytesIO()
    workbook.save(buffer)

    report = parse_normalized_price_workbook(buffer.getvalue())
    offer = next(item for item in report["candidates"] if item["product"] == "Gelato")

    assert offer["quoted_value"] == "€ 14,00"
    assert offer["numeric_price"] is None
    assert offer["status"] == "blocked"
    assert "package_price_not_single_unit" in offer["reasons"]


def test_unit_price_is_accepted_only_with_supplier_confirmation_status() -> None:
    workbook = Workbook()
    first = workbook.active
    first.title = "FOOD"
    for tab in EXPECTED_TABS:
        sheet = first if tab == "FOOD" else workbook.create_sheet(tab)
        sheet.append([f"{tab} — PREZZI ATTUALI"])
        sheet.append(["Prezzi attuali approvati"])
        sheet.append(["Aggiornamenti come nuove righe"])
        sheet.append(HEADERS)
    workbook["FOOD"].append([
        "Carni", "Gelato", "MELIUS", "Gelato 2,8 kg", "", "", "Confezione 2,8 kg",
        "€ 14,00", "2,8", "kg", "5,00",
        "Prezzo attuale approvato - prezzo unitario confermato dal fornitore",
        "Conferma fornitore", None, "mensile",
    ])
    buffer = BytesIO()
    workbook.save(buffer)

    report = parse_normalized_price_workbook(buffer.getvalue())
    offer = next(item for item in report["candidates"] if item["product"] == "Gelato")

    assert offer["quoted_value"] == "5,00"
    assert offer["numeric_price"] == "5.00"
    assert offer["source_cell"] == "K5"
    assert offer["status"] == "needs_mapping"


def test_normalized_workbook_rejects_missing_category_sheet() -> None:
    workbook = Workbook()
    workbook.active.title = "FOOD"
    buffer = BytesIO()
    workbook.save(buffer)
    try:
        parse_normalized_price_workbook(buffer.getvalue())
    except ValueError as error:
        assert "mancano le schede categoria" in str(error).lower()
    else:
        raise AssertionError("Missing category sheets must fail closed")


def test_new_workbook_imports_package_prices_without_units_and_ignores_old_supplier_ids() -> None:
    workbook = Workbook()
    workbook.active.title = "FOOD"
    for tab in NEW_FORMAT_TABS:
        sheet = workbook[tab] if tab == "FOOD" else workbook.create_sheet(tab)
        sheet.append(list(NEW_FORMAT_HEADERS.values()))

    def sku_for(source_key: str) -> str:
        return "SG3-" + hashlib.sha256(source_key.encode()).hexdigest()[:16].upper()

    workbook["FOOD"].append([
        sku_for("FOOD!2"), "GELATO", "Varie", "999", "MELIUS", "GELATO 2,8 KG",
        None, None, None, 14.20, None, "Origine FOOD!C2",
    ])
    workbook["FOOD"].append([
        sku_for("FOOD!2"), "GELATO", "Varie", "123", "MARR", None,
        None, None, None, 13.80, None, "Origine FOOD!D2",
    ])
    # Same description, different source product row: keep as a separate product.
    workbook["FOOD"].append([
        sku_for("FOOD!4"), "GELATO", "Varie", "999", "MELIUS", None,
        None, None, None, 12.50, None, "Origine FOOD!C4",
    ])
    workbook["BEVERAGE"].append([
        sku_for("BEVERAGE!2"), "ACQUA 50 CL", None, "999", "MELIUS", None,
        None, 1, "PZ", 0.35, None, "Origine BEVERAGE!C2",
    ])
    registry = workbook.create_sheet("FORNITORI")
    registry.append(["ID_FORNITORE", "FORNITORE", "PARTITA_IVA", "FOOD", "BEVERAGE", "FRUTTA_E_VERDURA", "MATERIALI", "NOTE"])
    registry.append(["999", "MELIUS", None, "SI", "SI", "NO", "NO", "ID vecchio da ignorare"])
    registry.append(["123", "MARR", None, "SI", "NO", "NO", "NO", ""])
    buffer = BytesIO()
    workbook.save(buffer)

    report = parse_normalized_price_workbook(buffer.getvalue())

    assert report["mode"] == "dry_run_only"
    assert report["candidate_count"] == 4
    assert report["source_supplier_sectors"] == {"MELIUS": ["Beverage", "Food"], "MARR": ["Food"]}
    gelato_offers = [item for item in report["candidates"] if item["product"] == "GELATO"]
    assert len(gelato_offers) == 3
    assert gelato_offers[0]["source_product_sku"] == sku_for("FOOD!2")
    assert gelato_offers[0]["source_product_sku"] == gelato_offers[1]["source_product_sku"]
    assert gelato_offers[0]["source_product_sku"] != gelato_offers[2]["source_product_sku"]
    assert gelato_offers[0]["numeric_price"] == "14.2"
    assert gelato_offers[0]["status"] == "needs_mapping"
    assert gelato_offers[0]["reference_unit"] is None
    assert "unit_not_specified" in gelato_offers[0]["reasons"]
    assert "999" not in [item["supplier"] for item in report["candidates"]]


def test_new_workbook_does_not_treat_multi_unit_package_amount_as_unit_price() -> None:
    workbook = Workbook()
    workbook.active.title = "FOOD"
    for tab in NEW_FORMAT_TABS:
        sheet = workbook[tab] if tab == "FOOD" else workbook.create_sheet(tab)
        sheet.append(list(NEW_FORMAT_HEADERS.values()))
    registry = workbook.create_sheet("FORNITORI")
    registry.append(["ID_FORNITORE", "FORNITORE", "PARTITA_IVA", "FOOD", "BEVERAGE", "FRUTTA_E_VERDURA", "MATERIALI", "NOTE"])
    key = "FOOD!2"
    sku = "SG3-" + hashlib.sha256(key.encode()).hexdigest()[:16].upper()
    workbook["FOOD"].append([sku, "GELATO 2,8 KG", "Varie", "old-id", "MELIUS", None, None, 2.8, "KG", 14.2, None, "Origine FOOD!C2"])
    buffer = BytesIO()
    workbook.save(buffer)

    offer = parse_normalized_price_workbook(buffer.getvalue())["candidates"][0]

    assert offer["numeric_price"] == "14.2"
    assert offer["reference_unit"] is None
    assert offer["supplier_price_unit"] is None
    assert "package_price_not_single_unit" in offer["reasons"]
