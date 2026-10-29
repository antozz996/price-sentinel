"""Read-only parser for the normalized PRICE SENTINEL category workbook."""
from __future__ import annotations

from collections import Counter
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
from io import BytesIO
import re
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


EXPECTED_TABS = {
    "FOOD": "Food",
    "BEVERAGE": "Beverage",
    "FRUTTA E VERDURA": "Food",
    "MATERIALI": "Materiali",
    "GIOCATTOLI": "Giocattoli",
    "ALTRO": "Altro",
}
NEW_FORMAT_TABS = {
    "FOOD": "Food",
    "BEVERAGE": "Beverage",
    "FRUTTA E VERDURA": "Food",
    "MATERIALI": "Materiali",
}
NEW_FORMAT_HEADERS = {
    "id_prodotto": "ID_PRODOTTO",
    "prodotto": "PRODOTTO",
    "sottocategoria": "SOTTOCATEGORIA",
    "id_fornitore": "ID_FORNITORE",
    "fornitore": "FORNITORE",
    "dicitura_fornitore": "DICITURA_FORNITORE",
    "marca": "MARCA",
    "quantita_confezione": "QUANTITA_CONFEZIONE",
    "unita_misura": "UNITA_MISURA",
    "prezzo_confezione_eur": "PREZZO_CONFEZIONE_EUR",
    "data_prezzo": "DATA_PREZZO",
    "note": "NOTE",
}
SUPPLIER_HEADERS = {
    "fornitore": "FORNITORE",
    "partita_iva": "PARTITA_IVA",
    "food": "FOOD",
    "beverage": "BEVERAGE",
    "frutta_e_verdura": "FRUTTA_E_VERDURA",
    "materiali": "MATERIALI",
}
REQUIRED_HEADERS = {
    "sottocategoria": "Sottocategoria",
    "prodotto di riferimento": "Prodotto di riferimento",
    "fornitore": "Fornitore",
    "dicitura del fornitore": "Dicitura del fornitore",
    "codice fornitore": "Codice fornitore",
    "marca": "Marca",
    "formato / confezione": "Formato / confezione",
    "prezzo confezione": "Prezzo confezione",
    "quantita confronto": "Quantità confronto",
    "unita confronto": "Unità confronto",
    "prezzo unitario": "Prezzo unitario",
    "stato": "Stato",
    "origine": "Origine",
    "data prezzo": "Data prezzo",
    "cadenza aggiornamento": "Cadenza aggiornamento",
}
APPROVED_STATUS = "prezzo attuale approvato"
SUPPLIER_CONFIRMED_UNIT_PRICE_STATUS = "prezzo attuale approvato - prezzo unitario confermato dal fornitore"
EXCLUDED_STATUS = "duplicata - non importare"


def _normalize(value: Any) -> str:
    text = str(value or "").strip().casefold()
    text = text.replace("à", "a").replace("á", "a").replace("è", "e").replace("é", "e")
    text = text.replace("ì", "i").replace("ò", "o").replace("ù", "u")
    return re.sub(r"\s+", " ", text)


def _text(value: Any) -> str:
    if value is None or isinstance(value, bool):
        return ""
    return str(value).strip()


def _parse_price(value: Any) -> tuple[str | None, str]:
    if value is None or isinstance(value, bool):
        return None, "missing"
    if isinstance(value, (int, float, Decimal)):
        raw = str(value)
    else:
        raw = _text(value)
        if not raw:
            return None, "missing"
        raw = re.sub(r"^(?:€|EUR)\s*|\s*(?:€|EUR)$", "", raw, flags=re.IGNORECASE).strip()
        raw = raw.replace(" ", "")
        # Accept unambiguous Italian decimal notation and grouped thousands.
        if re.fullmatch(r"\d{1,3}(?:\.\d{3})+,\d+", raw):
            raw = raw.replace(".", "").replace(",", ".")
        elif re.fullmatch(r"\d+(?:,\d+)?", raw):
            raw = raw.replace(",", ".")
        elif not re.fullmatch(r"\d+(?:\.\d+)?", raw):
            return None, "invalid_price"
    try:
        amount = Decimal(raw)
    except InvalidOperation:
        return None, "invalid_price"
    if not amount.is_finite() or amount <= 0:
        return None, "zero_or_negative"
    return format(amount, "f"), "valid"


def _parse_quantity(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool):
        return None
    raw = _text(value).replace(" ", "")
    if not raw:
        return None
    if re.fullmatch(r"\d+(?:,\d+)?", raw):
        raw = raw.replace(",", ".")
    try:
        amount = Decimal(raw)
    except InvalidOperation:
        return None
    return amount if amount.is_finite() and amount > 0 else None


def _header_map(row: tuple[Any, ...], required: dict[str, str]) -> dict[str, int] | None:
    indices = {_normalize(value).replace(" ", "_"): index for index, value in enumerate(row) if value is not None}
    if not set(required).issubset(indices):
        return None
    return indices


def _new_format_preview(workbook: Any, file_bytes: bytes) -> dict[str, Any]:
    missing = sorted((set(NEW_FORMAT_TABS) | {"FORNITORI"}) - set(workbook.sheetnames))
    if missing:
        raise ValueError("Nel nuovo file mancano le schede: " + ", ".join(missing))

    supplier_sheet = workbook["FORNITORI"]
    supplier_header = next((row for row in supplier_sheet.iter_rows(min_row=1, max_row=10, values_only=True)
                            if _header_map(row, SUPPLIER_HEADERS)), None)
    if supplier_header is None:
        raise ValueError("La scheda FORNITORI non contiene le intestazioni previste.")
    supplier_indices = _header_map(supplier_header, SUPPLIER_HEADERS) or {}
    source_supplier_sectors: dict[str, list[str]] = {}
    source_suppliers: list[str] = []
    sector_columns = {
        "food": "Food", "beverage": "Beverage",
        "frutta_e_verdura": "Food", "materiali": "Materiali",
    }
    for row in supplier_sheet.iter_rows(values_only=True):
        name = _text(row[supplier_indices["fornitore"]] if supplier_indices["fornitore"] < len(row) else None)
        if not name or _normalize(name) == "fornitore":
            continue
        sectors = [sector for column, sector in sector_columns.items()
                   if _normalize(row[supplier_indices[column]] if supplier_indices[column] < len(row) else None) in {"si", "sì", "yes", "true", "1"}]
        source_suppliers.append(name)
        source_supplier_sectors[name] = sorted(set(sectors))

    candidates: list[dict[str, Any]] = []
    catalog_items: dict[str, dict[str, Any]] = {}
    warnings: list[dict[str, Any]] = []
    sector_summary: list[dict[str, Any]] = []
    excluded_count = 0
    for tab, sector in NEW_FORMAT_TABS.items():
        sheet = workbook[tab]
        header_row = None
        header_row_number = 0
        indices = None
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            indices = _header_map(row, NEW_FORMAT_HEADERS)
            if indices is not None:
                header_row, header_row_number = row, row_number
                break
            if row_number >= 10:
                break
        if header_row is None or indices is None:
            raise ValueError(f"La scheda {tab} non contiene le intestazioni del nuovo formato.")
        tally = {"sector": sector, "tab": tab, "products": 0, "quotes": 0,
                 "excluded": 0, "blocked": 0, "review": 0, "skipped_empty": 0,
                 "duplicate_products": 0}
        product_ids: set[str] = set()
        for row_number, row in enumerate(sheet.iter_rows(min_row=header_row_number + 1, values_only=True), start=header_row_number + 1):
            if row_number - header_row_number > 5000:
                raise ValueError(f"La scheda {tab} supera il limite di 5.000 righe dati.")
            def cell(key: str) -> Any:
                index = indices[key]
                return row[index] if index < len(row) else None

            product_id = _text(cell("id_prodotto"))
            product = _text(cell("prodotto"))
            supplier = _text(cell("fornitore"))
            if not product_id and not product and not supplier:
                continue
            notes = _text(cell("note"))
            offer_id = _text(row[12]) if len(row) > 12 else ""
            if offer_id and (len(header_row) < 13 or _text(header_row[12]) != "ID_OFFERTA"):
                raise ValueError(f"La scheda {tab} contiene ID offerta senza l'intestazione ID_OFFERTA.")
            origin_match = re.search(r"Origine\s+(.+?!)(?:[A-Z]+)?([0-9]+)", notes)
            source_key = origin_match.group(1) + origin_match.group(2) if origin_match else f"{tab}!{row_number}"
            generated_sku = "SG3-" + hashlib.sha256(source_key.encode("utf-8")).hexdigest()[:16].upper()
            if not product_id:
                product_id = generated_sku
            elif product_id != generated_sku:
                warnings.append({"tab": tab, "row": row_number, "code": "product_id_provenance_mismatch", "id_prodotto": product_id})
                source_key = f"workbook-product:{product_id}"
                generated_sku = product_id

            unit = _text(cell("unita_misura"))
            quantity = _parse_quantity(cell("quantita_confezione"))
            canonical_unit = unit if unit and quantity == Decimal("1") else None
            catalog_item = {
                "source_sheet": tab,
                "source_product_key": source_key,
                "source_product_sku": product_id,
                "canonical_name": product,
                "comparison_unit": canonical_unit,
                "category": sector,
                "subcategory": _text(cell("sottocategoria")) or ("Frutta" if tab == "FRUTTA E VERDURA" else None),
            }
            if product and product_id:
                previous_item = catalog_items.get(product_id)
                if previous_item and _normalize(previous_item["canonical_name"]) != _normalize(product):
                    raise ValueError(f"L’ID prodotto {product_id} identifica nomi diversi nella scheda {tab}.")
                if previous_item:
                    if not previous_item["comparison_unit"] and canonical_unit:
                        previous_item["comparison_unit"] = canonical_unit
                    if not previous_item["subcategory"] and catalog_item["subcategory"]:
                        previous_item["subcategory"] = catalog_item["subcategory"]
                else:
                    catalog_items[product_id] = catalog_item
                if product_id not in product_ids:
                    tally["products"] += 1
                    product_ids.add(product_id)
            if supplier:
                source_supplier_sectors.setdefault(supplier, [])
                if supplier not in source_suppliers:
                    source_suppliers.append(supplier)
                if sector not in source_supplier_sectors[supplier]:
                    source_supplier_sectors[supplier].append(sector)

            if "esclusa nel santo graal" in _normalize(notes):
                excluded_count += 1
                tally["excluded"] += 1
                continue
            raw_price = cell("prezzo_confezione_eur")
            if not _text(raw_price):
                tally["skipped_empty"] += 1
                continue
            if not product_id or not product or not supplier:
                warnings.append({"tab": tab, "row": row_number, "code": "incomplete_offer_row"})
                tally["blocked"] += 1
                continue

            price, price_state = _parse_price(raw_price)
            supplier_confirmed = "conferma fornitore price sentinel" in _normalize(notes)
            complete_pack = supplier_confirmed and bool(unit and _text(cell("marca")) and _text(cell("dicitura_fornitore"))) and quantity is not None and quantity > 0
            reasons = ["canonical_id_not_mapped"]
            status = "needs_mapping"
            if price_state != "valid":
                status = "blocked"
                reasons.append(price_state)
            # The workbook amount is explicitly a package amount. Do not turn
            # it into a unit amount without a confirmed package quantity.
            if not unit:
                reasons.append("unit_not_specified")
            elif quantity and quantity != Decimal("1") and not complete_pack:
                reasons.append("package_price_not_single_unit")
            if price is not None and Decimal(price).as_tuple().exponent < -4:
                reasons.append("precision_review")

            display_price = format((Decimal(price) / quantity).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP), "f") if price is not None and complete_pack else price if price is not None else raw_price
            source_cell = f"{get_column_letter(indices['prezzo_confezione_eur'] + 1)}{row_number}"
            candidates.append({
                "source_product_key": source_key,
                "source_product_sku": product_id,
                "source_offer_id": offer_id or None,
                "supplier_description": _text(cell("dicitura_fornitore")) or product,
                "supplier_brand": _text(cell("marca")) or None,
                "package_quantity": format(quantity, "f") if quantity is not None else None,
                "package_unit": unit or None,
                "package_price": price,
                "package_confirmed": complete_pack,
                "sector": sector,
                "subcategory": _text(cell("sottocategoria")) or ("Frutta" if tab == "FRUTTA E VERDURA" else None),
                "source_sheet": tab,
                "source_row": row_number,
                "source_original_row": row_number,
                "source_cell": source_cell,
                "product": product,
                "supplier": supplier,
                "price_role": "negotiated_agreement",
                "requires_product_mapping": True,
                "reference_unit": unit if complete_pack or unit and quantity == Decimal("1") else None,
                "supplier_price_unit": unit if complete_pack or unit and quantity == Decimal("1") else None,
                "quoted_value": raw_price,
                "numeric_price": display_price,
                "status": status,
                "reasons": reasons,
                "row_note": _text(cell("dicitura_fornitore")) or None,
                "packaging_note": "; ".join(
                    f"{label}: {value}" for key, label in (
                        ("marca", "Marca"), ("quantita_confezione", "Quantità confezione"),
                        ("unita_misura", "Unità misura dichiarata"),
                    ) if (value := _text(cell(key)) )
                ) or "Prezzo importato come prezzo confezione; formato o unità da completare col fornitore.",
                "cell_note": notes or None,
            })
            tally["quotes"] += 1
            if status == "blocked":
                tally["blocked"] += 1
            else:
                tally["review"] += 1
        sector_summary.append(tally)

    offer_id_counts = Counter(item["source_offer_id"] for item in candidates if item.get("source_offer_id"))
    for item in candidates:
        if item.get("source_offer_id") and offer_id_counts[item["source_offer_id"]] > 1:
            if item["status"] != "blocked":
                next(entry for entry in sector_summary if entry["tab"] == item["source_sheet"])["blocked"] += 1
                next(entry for entry in sector_summary if entry["tab"] == item["source_sheet"])["review"] -= 1
            item["status"] = "blocked"
            item["reasons"] = sorted(set(item["reasons"] + ["duplicate_offer_id"]))

    return {
        "mode": "dry_run_only",
        "spreadsheet_id": "workbook:" + hashlib.sha256(file_bytes).hexdigest(),
        "source_title_expected": "PRICE SENTINEL — SANTO GRAAL",
        "workbook_title": "PRICE SENTINEL — SANTO GRAAL",
        "no_contract_writes": True,
        "candidate_count": len(candidates),
        "skipped_empty_count": sum(item["skipped_empty"] for item in sector_summary),
        "sector_summary": sector_summary,
        "catalog_items": list(catalog_items.values()),
        "food_audit": None,
        "missing_sources": [],
        "warnings": warnings,
        "source_suppliers": list(dict.fromkeys(source_suppliers)),
        "source_supplier_sectors": source_supplier_sectors,
        "candidates": candidates,
        "rules": {
            "agreement": "valid_until_next_approved_listino_unless_explicit_expiry",
            "auto_commit_enabled": False,
            "canonical_product_mapping": "required_before_import",
            "source": "Nuovo workbook Santo Graal, importo della confezione come dichiarato.",
            "price_semantics": "package_price; never converted to unit price without confirmed package quantity",
        },
        "excluded_count": excluded_count,
    }


def parse_normalized_sheet_values(values: dict[str, list[list[Any]]], spreadsheet_id: str) -> dict[str, Any]:
    """Parse a Google Sheets snapshot with the same rules as the uploaded workbook."""
    expected = set(NEW_FORMAT_TABS) | {"FORNITORI"}
    if set(values) != expected:
        raise ValueError("Le schede del nuovo Sheet non corrispondono al formato previsto.")
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name in (*NEW_FORMAT_TABS, "FORNITORI"):
        rows = values[name]
        if len(rows) > 5001:
            raise ValueError(f"La scheda {name} supera il limite di 5.000 righe dati.")
        sheet = workbook.create_sheet(name)
        for row in rows:
            if len(row) > (8 if name == "FORNITORI" else 13):
                raise ValueError(f"La scheda {name} contiene colonne inattese.")
            if any(isinstance(value, str) and value.startswith("=") for value in row):
                raise ValueError(f"La scheda {name} contiene formule: usa valori calcolati per l’anteprima.")
            sheet.append(list(row))
    output = BytesIO()
    workbook.save(output)
    report = parse_normalized_price_workbook(output.getvalue())
    report["spreadsheet_id"] = spreadsheet_id
    return report


def parse_normalized_price_workbook(file_bytes: bytes) -> dict[str, Any]:
    """Return the same candidate shape as the Google Sheet dry-run, without DB writes."""
    try:
        workbook = load_workbook(BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("File Excel non leggibile. Usa il workbook migrato in formato .xlsx.") from exc

    try:
        if set(NEW_FORMAT_TABS).issubset(workbook.sheetnames) and "FORNITORI" in workbook.sheetnames:
            return _new_format_preview(workbook, file_bytes)
        missing = sorted(set(EXPECTED_TABS) - set(workbook.sheetnames))
        if missing:
            raise ValueError("Nel file mancano le schede categoria: " + ", ".join(missing))

        candidates: list[dict[str, Any]] = []
        warnings: list[dict[str, Any]] = []
        sector_summary: list[dict[str, Any]] = []
        duplicate_keys: Counter[tuple[str, str, str]] = Counter()
        staged_rows: list[dict[str, Any]] = []
        excluded_count = 0

        for tab, sector in EXPECTED_TABS.items():
            sheet = workbook[tab]
            iterator = sheet.iter_rows(values_only=True)
            header: list[Any] | None = None
            header_row_number = 0
            for row_number, row in enumerate(iterator, start=1):
                normalized_headers = {_normalize(value): index for index, value in enumerate(row) if value is not None}
                if set(REQUIRED_HEADERS).issubset(normalized_headers):
                    header = list(row)
                    header_row_number = row_number
                    break
                if row_number >= 15:
                    break
            if header is None:
                raise ValueError(f"La scheda {tab} non contiene l’intestazione prevista del file migrato.")
            indices = {_normalize(value): index for index, value in enumerate(header) if value is not None}
            tally = {"sector": sector, "tab": tab, "products": 0, "quotes": 0,
                     "excluded": 0, "blocked": 0, "review": 0, "skipped_empty": 0,
                     "duplicate_products": 0}
            seen_product_names: set[str] = set()
            for row_number, row in enumerate(iterator, start=header_row_number + 1):
                if row_number - header_row_number > 5000:
                    raise ValueError(f"La scheda {tab} supera il limite di 5.000 righe dati.")
                product = _text(row[indices["prodotto di riferimento"]] if indices["prodotto di riferimento"] < len(row) else None)
                supplier = _text(row[indices["fornitore"]] if indices["fornitore"] < len(row) else None)
                original_description = _text(row[indices["dicitura del fornitore"]] if indices["dicitura del fornitore"] < len(row) else None)
                # The migrated reference column may contain unapproved grouping
                # (e.g. BASILICO). Preserve each original Sheet product wording.
                if original_description:
                    product = original_description
                if not product and not supplier:
                    continue
                state = _normalize(row[indices["stato"]] if indices["stato"] < len(row) else None)
                if state == EXCLUDED_STATUS:
                    excluded_count += 1
                    tally["excluded"] += 1
                    continue
                package_price = row[indices["prezzo confezione"]] if indices["prezzo confezione"] < len(row) else None
                explicit_unit_price = row[indices["prezzo unitario"]] if indices["prezzo unitario"] < len(row) else None
                unit_price_confirmed = state == SUPPLIER_CONFIRMED_UNIT_PRICE_STATUS and bool(_text(explicit_unit_price))
                if not _text(package_price) and not unit_price_confirmed:
                    # Empty future products are intentionally left for later entry.
                    tally["skipped_empty"] += 1
                    continue
                if not product or not supplier:
                    warnings.append({"tab": tab, "row": row_number, "code": "incomplete_offer_row"})
                    continue

                tally["products"] += int(_normalize(product) not in seen_product_names)
                seen_product_names.add(_normalize(product))
                base_unit = _text(row[indices["unita confronto"]] if indices["unita confronto"] < len(row) else None)
                raw_quantity = row[indices["quantita confronto"]] if indices["quantita confronto"] < len(row) else None
                quantity = _parse_quantity(raw_quantity)
                raw_price = explicit_unit_price if unit_price_confirmed else package_price
                price, price_state = _parse_price(raw_price)
                reasons: list[str] = ["canonical_id_not_mapped"]
                status = "needs_mapping"
                if state not in {APPROVED_STATUS, SUPPLIER_CONFIRMED_UNIT_PRICE_STATUS}:
                    status = "blocked"
                    reasons.append("not_current_approved")
                if price_state != "valid":
                    status = "blocked"
                    reasons.append(price_state)
                if not unit_price_confirmed and quantity != Decimal("1"):
                    status = "blocked"
                    reasons.append("package_price_not_single_unit")
                if not base_unit:
                    status = "blocked"
                    reasons.append("unit_not_specified")
                elif _normalize(base_unit) == "altro":
                    status = "blocked"
                    reasons.append("unit_mapping_needed")
                if price is not None and Decimal(price).as_tuple().exponent < -4:
                    reasons.append("precision_review")

                subcategory = _text(row[indices["sottocategoria"]] if indices["sottocategoria"] < len(row) else None) or None
                if tab == "FRUTTA E VERDURA" and not subcategory:
                    subcategory = "Frutta"
                offer = {
                    "source_product_key": re.sub(r"![A-Z]+([0-9]+)$", r"!\1", _text(row[indices["origine"]] if indices["origine"] < len(row) else None)) or f"{tab}!{row_number}",
                    "sector": sector,
                    "subcategory": subcategory,
                    "source_sheet": tab,
                    "source_row": row_number,
                    "source_original_row": row_number,
                    "source_cell": f"{get_column_letter((indices['prezzo unitario'] if unit_price_confirmed else indices['prezzo confezione']) + 1)}{row_number}",
                    "product": product,
                    "supplier": supplier,
                    "price_role": "negotiated_agreement",
                    "requires_product_mapping": True,
                    "reference_unit": base_unit or None,
                    "supplier_price_unit": base_unit if unit_price_confirmed or quantity == Decimal("1") else None,
                    "quoted_value": raw_price,
                    "numeric_price": price if (unit_price_confirmed or quantity == Decimal("1")) and base_unit else None,
                    "status": status,
                    "reasons": reasons,
                    "row_note": _text(row[indices["dicitura del fornitore"]] if indices["dicitura del fornitore"] < len(row) else None) or None,
                    "packaging_note": "; ".join(
                        f"{label}: {value}" for key, label in (
                            ("codice fornitore", "Codice fornitore"), ("marca", "Marca"),
                            ("formato / confezione", "Formato / confezione"),
                            ("prezzo confezione", "Prezzo confezione (dato fornitore)"),
                            ("prezzo unitario", "Prezzo unitario confermato"),
                            ("quantita confronto", "Quantità confronto dichiarata"),
                            ("unita confronto", "Unità confronto dichiarata"),
                        )
                        if (value := _text(row[indices[key]] if indices[key] < len(row) else None))
                    ) or None,
                    "cell_note": _text(row[indices["origine"]] if indices["origine"] < len(row) else None) or None,
                }
                if state not in {APPROVED_STATUS, SUPPLIER_CONFIRMED_UNIT_PRICE_STATUS}:
                    reasons.append("not_current_approved")
                duplicate_key = (tab, _normalize(product), _normalize(supplier))
                duplicate_keys[duplicate_key] += 1
                staged_rows.append(offer)
                tally["quotes"] += 1
            sector_summary.append(tally)

        duplicate_offer_count = 0
        for offer in staged_rows:
            offer['source_product_sku'] = 'SG3-' + hashlib.sha256(offer['source_product_key'].encode()).hexdigest()[:16].upper()
            duplicate_key = (offer["source_sheet"], _normalize(offer["product"]), _normalize(offer["supplier"]))
            if duplicate_keys[duplicate_key] > 1:
                offer["status"] = "blocked"
                offer["reasons"] = sorted(set(offer["reasons"] + ["duplicate_supplier_product_offer"]))
                duplicate_offer_count += 1
            if offer["status"] == "blocked":
                next(item for item in sector_summary if item["tab"] == offer["source_sheet"])["blocked"] += 1
            else:
                next(item for item in sector_summary if item["tab"] == offer["source_sheet"])["review"] += 1
            candidates.append(offer)

        if len(candidates) > 5000:
            raise ValueError("Il workbook supera il limite di 5.000 offerte per un’unica anteprima.")

        if excluded_count:
            warnings.append({"code": "explicitly_excluded_rows", "count": excluded_count})
        if duplicate_offer_count:
            warnings.append({"code": "duplicate_supplier_product_offer", "count": duplicate_offer_count})
        return {
            "mode": "dry_run_only",
            "spreadsheet_id": "workbook:" + hashlib.sha256(file_bytes).hexdigest(),
            "source_title_expected": "PRICE SENTINEL — SANTO GRAAL migrato",
            "workbook_title": "PRICE SENTINEL — SANTO GRAAL migrato",
            "no_contract_writes": True,
            "candidate_count": len(candidates),
            "skipped_empty_count": sum(item["skipped_empty"] for item in sector_summary),
            "sector_summary": sector_summary,
            "food_audit": None,
            "missing_sources": [],
            "warnings": warnings,
            "candidates": candidates,
            "rules": {
                "agreement": "valid_until_next_approved_listino_unless_explicit_expiry",
                "auto_commit_enabled": False,
                "canonical_product_mapping": "required_before_import",
                "source": "Prezzi correnti approvati del workbook migrato.",
                "price_semantics": "supplier_confirmed_unit_price_status_or_package_price_at_declared_quantity_one",
            },
            "excluded_count": excluded_count,
        }
    finally:
        workbook.close()
